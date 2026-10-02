#!/usr/bin/env python3
"""Build an EDITABLE, integrant-branded PowerPoint from deepeval-session.md.

Unlike the Marp PPTX export (one flat image per slide), this produces native
PowerPoint text boxes, bullet lists, code blocks and tables that can be edited
directly in PowerPoint. Branding (integrant logo, top/bottom rules, orange
corner square) is drawn on every slide to match the company template.

Usage:
    python build_integrant_pptx.py
"""
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
SRC_MD = HERE / "deepeval-session.md"
TEMPLATE = HERE / "Integrant-Basic-Presentation-En-Template.pptx"
LOGO = HERE / "assets" / "integrant-logo.png"
OUT = HERE / "deepeval-session-integrant-editable.pptx"

# ---- brand palette -------------------------------------------------------
CYAN      = RGBColor(0x20, 0xB9, 0xEC)
CYAN_DK   = RGBColor(0x13, 0x98, 0xC4)
ORANGE    = RGBColor(0xED, 0x7D, 0x31)
GRAY_DK   = RGBColor(0x44, 0x54, 0x6A)
GRAY      = RGBColor(0x6E, 0x77, 0x81)
TEXT      = RGBColor(0x24, 0x29, 0x2F)
CODE_TXT  = RGBColor(0x1F, 0x23, 0x28)
CODE_BG   = RGBColor(0xF6, 0xF8, 0xFA)
BORDER    = RGBColor(0xD0, 0xD7, 0xDE)
TH_TXT    = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT   = RGBColor(0xF3, 0xFB, 0xFE)
QUOTE_BG  = RGBColor(0xFD, 0xF3, 0xEA)
RULE_CLR  = RGBColor(0xBF, 0xBF, 0xBF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

TITLE_FONT = "Sofia Pro Medium"
BODY_FONT  = "Roboto"
CODE_FONT  = "Consolas"

EMU_IN = 914400
# content frame
CL = Inches(0.92)      # content left
CW = Inches(11.5)      # content width
TOP_RULE_Y = Inches(1.35)
BOT_RULE_Y = Inches(6.84)
BODY_TOP = Inches(1.62)
BODY_BOTTOM = Inches(6.66)


# ========================================================================
# Markdown parsing
# ========================================================================
def load_slides(md_text):
    # strip YAML front-matter (first --- ... ---)
    if md_text.startswith("---"):
        parts = md_text.split("\n---\n", 1)
        # front matter ends at the first '\n---\n' after the opening '---'
        body = md_text[3:]
        idx = body.find("\n---\n")
        md_text = body[idx + 5:]
    raw_slides = re.split(r"\n---\n", md_text)
    return [s.strip("\n") for s in raw_slides if s.strip()]


def extract_notes(block):
    """Return (clean_block, notes_text). Pull out <!-- ... --> that are not _class."""
    notes = []

    def repl(m):
        inner = m.group(1).strip()
        if inner.startswith("_class:"):
            return m.group(0)  # keep class directive for detection
        notes.append(inner)
        return ""

    clean = re.sub(r"<!--(.*?)-->", repl, block, flags=re.DOTALL)
    return clean, "\n\n".join(notes).strip()


def is_lead(block):
    return "_class: lead" in block


def strip_html(s):
    s = s.replace("<br>", "\n")
    s = re.sub(r"</?span[^>]*>", "", s)
    s = re.sub(r"</?div[^>]*>", "", s)
    return s


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def inline_runs(text):
    """Split text into runs: list of dicts {text, bold, italic, code}."""
    text = strip_html(text)
    runs = []
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            runs.append({"text": text[pos:m.start()]})
        tok = m.group(0)
        if tok.startswith("**"):
            runs.append({"text": tok[2:-2], "bold": True})
        elif tok.startswith("`"):
            runs.append({"text": tok[1:-1], "code": True})
        else:
            runs.append({"text": tok[1:-1], "italic": True})
        pos = m.end()
    if pos < len(text):
        runs.append({"text": text[pos:]})
    return runs or [{"text": ""}]


def parse_body(lines):
    """Parse content lines into ordered blocks."""
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        # fenced code
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # closing fence
            blocks.append({"type": "code", "lang": lang, "lines": buf})
            continue
        # table
        if stripped.startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            blocks.append({"type": "table", "rows": parse_table(tbl)})
            continue
        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)),
                           "text": m.group(2)})
            i += 1
            continue
        # blockquote (may span multiple > lines)
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append({"type": "quote", "text": " ".join(b.strip() for b in buf)})
            continue
        # list item
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            buf = []
            while i < n:
                mm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if mm:
                    indent = len(mm.group(1))
                    ordered = bool(re.match(r"\d+\.", mm.group(2)))
                    buf.append({"level": indent // 2, "ordered": ordered,
                                "text": mm.group(3)})
                    i += 1
                elif buf and lines[i].startswith("  ") and lines[i].strip() \
                        and not lines[i].strip().startswith(("```", "|", ">", "#")):
                    # lazy continuation line of the previous bullet
                    buf[-1]["text"] += " " + lines[i].strip()
                    i += 1
                elif lines[i].strip() == "":
                    # peek: continue list only if next non-blank is a list item
                    j = i + 1
                    while j < n and lines[j].strip() == "":
                        j += 1
                    if j < n and re.match(r"^(\s*)([-*]|\d+\.)\s+", lines[j]):
                        i += 1
                    else:
                        break
                else:
                    break
            blocks.append({"type": "list", "items": buf})
            continue
        # paragraph
        blocks.append({"type": "para", "text": stripped})
        i += 1
    return blocks


def parse_table(tbl_lines):
    rows = []
    for ln in tbl_lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    # drop the separator row (---)
    rows = [r for r in rows if not all(re.match(r"^:?-{2,}:?$", c or "-") for c in r)]
    return rows


# ========================================================================
# Rendering helpers
# ========================================================================
def set_runs(paragraph, runs, base_size, base_color, font=BODY_FONT):
    for r in runs:
        run = paragraph.add_run()
        run.text = r["text"]
        f = run.font
        f.size = base_size
        f.name = CODE_FONT if r.get("code") else font
        f.bold = bool(r.get("bold"))
        f.italic = bool(r.get("italic"))
        if r.get("code"):
            f.color.rgb = CYAN_DK
        elif r.get("bold"):
            f.color.rgb = ORANGE
        elif r.get("italic"):
            f.color.rgb = GRAY
        else:
            f.color.rgb = base_color


def add_branding(slide, page_no, top_rule=True):
    shapes = slide.shapes
    # bottom rule
    ln = shapes.add_connector(2, Inches(0.78), BOT_RULE_Y, Inches(12.51), BOT_RULE_Y)
    ln.line.color.rgb = RULE_CLR
    ln.line.width = Pt(1)
    # top rule below the title — part of the integrant DNA on every slide
    ln2 = shapes.add_connector(2, Inches(0.78), TOP_RULE_Y, Inches(12.51), TOP_RULE_Y)
    ln2.line.color.rgb = RULE_CLR
    ln2.line.width = Pt(1)
    # orange corner square
    sq = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.10), Inches(6.84),
                          Inches(0.41), Inches(0.41))
    sq.fill.solid()
    sq.fill.fore_color.rgb = ORANGE
    sq.line.fill.background()
    sq.shadow.inherit = False
    tf = sq.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(page_no)
    run.font.size = Pt(12)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.name = BODY_FONT
    # logo
    if LOGO.exists():
        shapes.add_picture(str(LOGO), Inches(0.82), Inches(6.90), height=Inches(0.35))


def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    return tb, tf


# rough char-per-line for wrap estimation at given font size over width inches
def est_lines(text, width_in, size_pt):
    cpl = max(10, int(width_in * 72 / (size_pt * 0.52)))
    total = 0
    for ln in text.split("\n"):
        total += max(1, -(-len(ln) // cpl))  # ceil
    return total


# ========================================================================
# Slide builders
# ========================================================================
def build_lead(slide, blocks):
    # gather headings/paras
    h1 = next((b["text"] for b in blocks if b["type"] == "heading" and b["level"] == 1), None)
    h2 = next((b["text"] for b in blocks if b["type"] == "heading" and b["level"] == 2), None)
    h3 = next((b["text"] for b in blocks if b["type"] == "heading" and b["level"] == 3), None)
    paras = [b["text"] for b in blocks if b["type"] == "para"]

    tb, tf = add_textbox(slide, Inches(1.0), Inches(2.4), Inches(11.33), Inches(2.6))
    tf.anchor = MSO_ANCHOR.MIDDLE
    first = True
    def para(text, size, color, bold=True, font=TITLE_FONT, space=10):
        nonlocal first
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(space)
        set_runs(p, inline_runs(text), Pt(size), color, font=font)
        for r in p.runs:
            r.font.bold = bold
        return p
    if h1:
        para(h1, 40, CYAN, bold=True)
    if h2:
        para(h2, 26, GRAY_DK, bold=True)
    if h3:
        para(h3, 20, GRAY, bold=False)
    for pr in paras:
        para(pr, 16, GRAY, bold=False, font=BODY_FONT)


def build_agenda(slide, blocks):
    """Replicate the integrant 'AGENDA' design from template slide 6:
    big AGENDA title, cyan numbered circles, Roboto item text beside each."""
    # title
    tb, tf = add_textbox(slide, Inches(0.84), Inches(0.50), Inches(4.2), Inches(0.80))
    tf.anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "AGENDA"
    run.font.size = Pt(40)
    run.font.name = TITLE_FONT
    run.font.color.rgb = CYAN

    # gather ordered items
    items = []
    for b in blocks:
        if b["type"] == "list":
            items.extend(b["items"])
    tops = [1.66, 2.31, 2.94, 3.57, 4.20, 4.83, 5.46, 6.09]
    n = len(items)
    if n > len(tops):
        # compress spacing to fit
        tops = [1.66 + i * (4.43 / max(1, n - 1)) for i in range(n)]
    for i, it in enumerate(items):
        y = tops[i] if i < len(tops) else 1.66 + i * 0.63
        # cyan numbered circle
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.84), Inches(y),
                                      Inches(0.47), Inches(0.47))
        oval.fill.solid()
        oval.fill.fore_color.rgb = CYAN
        oval.line.fill.background()
        oval.shadow.inherit = False
        otf = oval.text_frame
        otf.margin_left = 0; otf.margin_right = 0
        otf.margin_top = 0; otf.margin_bottom = 0
        op = otf.paragraphs[0]
        op.alignment = PP_ALIGN.CENTER
        orun = op.add_run()
        orun.text = str(i + 1)
        orun.font.size = Pt(16)
        orun.font.bold = True
        orun.font.name = BODY_FONT
        orun.font.color.rgb = WHITE
        # item text
        tb2, tf2 = add_textbox(slide, Inches(1.36), Inches(y + 0.02),
                               Inches(11.14), Inches(0.44))
        tf2.anchor = MSO_ANCHOR.MIDDLE
        pp = tf2.paragraphs[0]
        for r in inline_runs(it["text"]):
            rr = pp.add_run()
            rr.text = r["text"]
            rr.font.size = Pt(17)
            rr.font.name = CODE_FONT if r.get("code") else BODY_FONT
            if r.get("code"):
                rr.font.color.rgb = CYAN_DK
            elif r.get("bold"):
                rr.font.bold = True
                rr.font.color.rgb = TEXT
            else:
                rr.font.color.rgb = GRAY_DK


def build_content(slide, blocks):
    # title = first heading (any level); render remaining blocks in flow
    title = None
    start = 0
    for k, b in enumerate(blocks):
        if b["type"] == "heading":
            title = b
            start = k + 1
            break
    if title:
        tb, tf = add_textbox(slide, CL, Inches(0.46), CW, Inches(0.95))
        tf.anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        set_runs(p, inline_runs(title["text"]), Pt(28), CYAN, font=TITLE_FONT)
        for r in p.runs:
            r.font.bold = True
            if not r.font.color.rgb:
                r.font.color.rgb = CYAN

    body = blocks[start:]
    # estimate total height to pick a scale
    cursor = BODY_TOP
    avail = BODY_BOTTOM - BODY_TOP
    # first pass measure
    measured = [measure_block(b) for b in body]
    total = sum(h for h, _ in measured) + Inches(0.08) * max(0, len(body) - 1)
    scale = 1.0
    if total > avail and total > 0:
        scale = max(0.72, avail / total)
    for b, (h, _) in zip(body, measured):
        hh = Emu(int(h * scale))
        render_block(slide, b, cursor, hh, scale)
        cursor = Emu(cursor + hh + Inches(0.08))


FS_BODY = 15
FS_CODE = 11.5
FS_TABLE = 12


def measure_block(b):
    if b["type"] == "code":
        nlines = max(1, len(b["lines"]))
        h = Inches(0.26 * (FS_CODE / 11.5)) * nlines + Inches(0.22)
        return int(h), b
    if b["type"] == "table":
        rows = len(b["rows"])
        return int(Inches(0.40) * rows + Inches(0.05)), b
    if b["type"] == "list":
        lines = 0
        for it in b["items"]:
            lines += est_lines(strip_html(it["text"]), 11.0, FS_BODY)
        return int(Inches(0.30) * lines + Inches(0.06)), b
    if b["type"] == "quote":
        lines = est_lines(strip_html(b["text"]), 10.8, FS_BODY)
        return int(Inches(0.30) * lines + Inches(0.18)), b
    if b["type"] == "heading":
        return int(Inches(0.42)), b
    # para
    lines = est_lines(strip_html(b["text"]), 11.2, FS_BODY)
    return int(Inches(0.30) * lines + Inches(0.06)), b


def render_block(slide, b, top, height, scale):
    if b["type"] == "code":
        render_code(slide, b, top, height, scale)
    elif b["type"] == "table":
        render_table(slide, b, top, height, scale)
    elif b["type"] == "list":
        render_list(slide, b, top, height, scale)
    elif b["type"] == "quote":
        render_quote(slide, b, top, height, scale)
    elif b["type"] == "heading":
        render_subheading(slide, b, top, height, scale)
    else:
        render_para(slide, b, top, height, scale)


def render_subheading(slide, b, top, height, scale):
    tb, tf = add_textbox(slide, CL, top, CW, height)
    p = tf.paragraphs[0]
    size = {1: 26, 2: 22, 3: 18}.get(b["level"], 18)
    set_runs(p, inline_runs(b["text"]), Pt(size * scale),
             CYAN_DK if b["level"] <= 2 else GRAY_DK, font=TITLE_FONT)
    for r in p.runs:
        r.font.bold = True


def render_para(slide, b, top, height, scale):
    tb, tf = add_textbox(slide, CL, top, CW, height)
    p = tf.paragraphs[0]
    set_runs(p, inline_runs(b["text"]), Pt(FS_BODY * scale), TEXT)


def render_list(slide, b, top, height, scale):
    tb, tf = add_textbox(slide, CL, top, CW, height)
    first = True
    for it in b["items"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = min(4, it["level"])
        p.space_after = Pt(3)
        bullet = "• " if not it["ordered"] else ""
        indent = "   " * it["level"]
        if bullet:
            r0 = p.add_run()
            r0.text = indent + bullet
            r0.font.size = Pt(FS_BODY * scale)
            r0.font.name = BODY_FONT
            r0.font.color.rgb = CYAN
            r0.font.bold = True
        set_runs(p, inline_runs(it["text"]), Pt(FS_BODY * scale), TEXT)


def render_quote(slide, b, top, height, scale):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, CL, top, CW, height)
    box.fill.solid()
    box.fill.fore_color.rgb = QUOTE_BG
    box.line.fill.background()
    box.shadow.inherit = False
    # left accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, CL, top, Inches(0.06), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    bar.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    set_runs(p, inline_runs(b["text"]), Pt(FS_BODY * scale), GRAY_DK)
    for r in p.runs:
        if not (r.font.bold or r.font.italic):
            r.font.color.rgb = GRAY_DK


def render_code(slide, b, top, height, scale):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, CL, top, CW, height)
    box.adjustments[0] = 0.04
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = BORDER
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    first = True
    for ln in (b["lines"] or [""]):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.0
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = ln if ln else " "
        run.font.name = CODE_FONT
        run.font.size = Pt(FS_CODE * scale)
        run.font.color.rgb = CODE_TXT


def render_table(slide, b, top, height, scale):
    rows = b["rows"]
    nrows = len(rows)
    ncols = max(len(r) for r in rows)
    gfx = slide.shapes.add_table(nrows, ncols, CL, top, CW, height)
    table = gfx.table
    table.first_row = False
    table.horz_banding = False
    for ci in range(ncols):
        pass
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            text = row[ci] if ci < len(row) else ""
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = CYAN
            elif ri % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = ROW_ALT
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            tfc = cell.text_frame
            tfc.word_wrap = True
            p = tfc.paragraphs[0]
            if ri == 0:
                run = p.add_run(); run.text = strip_html(text).replace("**", "")
                run.font.bold = True
                run.font.size = Pt(FS_TABLE * scale)
                run.font.color.rgb = TH_TXT
                run.font.name = BODY_FONT
            else:
                set_runs(p, inline_runs(text), Pt(FS_TABLE * scale), TEXT)
    # row heights
    rh = int(height / nrows)
    for r in table.rows:
        r.height = rh


# ========================================================================
# Main
# ========================================================================
def remove_all_slides(prs):
    xml_slides = prs.slides._sldIdLst
    for sldId in list(xml_slides):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        xml_slides.remove(sldId)


def main():
    md = SRC_MD.read_text(encoding="utf-8")
    slides = load_slides(md)

    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    blank = prs.slide_layouts[6]  # 'Blank'

    page = 0
    for raw in slides:
        clean, notes = extract_notes(raw)
        lead = is_lead(clean)
        # remove class directive comments now
        clean = re.sub(r"<!--.*?-->", "", clean, flags=re.DOTALL)
        lines = clean.split("\n")
        blocks = parse_body(lines)
        if not blocks:
            continue
        page += 1
        slide = prs.slides.add_slide(blank)
        title_block = next((b for b in blocks if b["type"] == "heading"), None)
        is_agenda = bool(title_block) and title_block["text"].strip().lower() == "agenda"
        add_branding(slide, page, top_rule=(lead or is_agenda))
        if lead:
            build_lead(slide, blocks)
        elif is_agenda:
            build_agenda(slide, blocks)
        else:
            build_content(slide, blocks)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(str(OUT))
    print(f"wrote {OUT}  ({page} slides)")


if __name__ == "__main__":
    main()
